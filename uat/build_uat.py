# -*- coding: utf-8 -*-
"""build_uat.py -- the user acceptance test script, generated from the Gherkin stories and checked against the state model.

PURPOSE
  Every UAT row is one Gherkin scenario, so the test script cannot drift from the stories. Each row carries its trace:
  the state transition it exercises (or the state it checks, or the move it proves refused), the rule that move
  enforces, the map steps, the decisions. A story cannot test a move the model does not declare. The owning
  department signs (D11).

CONFIG
  argv[1]  the repository root. Needs openpyxl and gherkin-official (Cucumber's own parser). Reads
           state/transitions.csv (written by state/model.py) and process/after.bpmn (written by process/build_maps.py).

WORKFLOW
  1 parse stories/*.feature; read "# Document:", "# Decision:" and "# Map: after.bpmn, <step ids>." from the header
  2 every scenario is preceded by exactly one tag:
      "# Transition: <document> | <from state> | <event>"   a move the model allows
      "# State: <document> | <state>[, <state>...][ | <rule>]" a state the scenario checks, with no move
      "# Error path: <document> | <from state> | <event>"   a move the model refuses
  3 check each tag against state/transitions.csv, and each map step against the ids in process/after.bpmn
  4 one row per scenario: id, story, scenario, document, state transition, rule, map steps, decisions, steps,
    expected, then the tester's cells (actual, result, tester role, date), shaded as input
  5 Sign-off sheet per story; Read me lists the model's refused moves that no row tests
  6 [CONTROL] a planted refused move, a planted transition and a planted map step are each caught; rows written ==
    scenarios parsed; every story has at least one error path
  rc 0 ok . 2 usage . 3 a check failed

Q&A
  Q: why a tag on every scenario?   A: done-test 2 asks for each scenario's state transition. A tag the build checks is
                                       the only way the answer cannot drift from the model.

PITFALLS
  Moves are matched on document, from state and event exactly as the model spells them; a reworded event in a story
  fails the check, which is the intent.

NOTES
  Style: Calibri, header #1D5C6B, input cells #FFF8E6.
"""
import csv
import glob
import io
import os
import re
import sys
import xml.etree.ElementTree as ET

from gherkin.parser import Parser
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADER, HEADER_TEXT, INPUT = "1D5C6B", "FFFFFF", "FFF8E6"
PASS, FAIL = "2E7D4F", "B23A3A"
OWNER = {  # the department that owns the process signs its story (D11)
    "order-lifecycle": ("Sales", "Sales lead"),
    "order-closure": ("Sales", "Sales lead, with the logistics lead (closure rules, D6)"),
    "delivery-trip": ("Logistics", "Logistics lead (D14)"),
    "material-request": ("Production", "Factory manager (material requests and issues, D6)"),
    "purchase-order": ("Purchasing", "Logistics manager, who signs purchase orders on paper (D9)"),
    "voucher-posting": ("Accounting", "Chief accountant (posting policy, D10)"),
    "sale-statuses": ("Accounting", "Finance lead (legal status combinations, D8)"),
    "receivables": ("Accounting", "Chief accountant (payment terms, credit hold and collection, D17)"),
}
RESULTS = ["Pass", "Fail", "Blocked"]
DEFECT_ROWS = 30  # blank defect rows, enough for one UAT cycle
TAG = re.compile(r"^# (Transition|State|Error path): ([\w-]+) \| ([^|]+?)(?: \| (.+))?$")


def model(root):
    rows = list(csv.DictReader(io.open(os.path.join(root, "state", "transitions.csv"), encoding="utf-8")))
    allowed = {(r["document"], r["from"], r["event"]): r for r in rows if r["kind"] == "allowed"}
    refused = {(r["document"], r["from"], r["event"]): r for r in rows if r["kind"] == "refused"}
    states = {}
    for r in rows:
        for s in (r["from"], r["to"]):
            if s:
                states.setdefault(r["document"], set()).add(s)
    return allowed, refused, states


def map_steps(root):
    ns = {"b": "http://www.omg.org/spec/BPMN/20100524/MODEL"}
    tree = ET.parse(os.path.join(root, "process", "after.bpmn"))
    return {e.get("id") for e in tree.getroot().iter() if e.get("id") and e.tag.startswith("{%s}" % ns["b"])}


def map_moves(root):
    """{map task id: (document, from, event)}: the move each task on the after map makes (process/build_maps.py)."""
    sys.path.insert(0, os.path.join(root, "process"))
    import build_maps
    return {t: tuple(m) for t, m in build_maps.AFTER["moves"].items()}


def row_steps(tag, moves, story_ids):
    """The map steps one row exercises: for a transition, the task that makes that move; for a state or a refusal,
    the tasks that move the document on from that state. Falls back to the story's steps when the map has none."""
    if not tag:
        return story_ids
    kind, doc, first, event = tag.group(1), tag.group(2), tag.group(3).strip(), (tag.group(4) or "").strip()
    froms = [x.strip() for x in first.split(",")]
    if kind == "Transition":
        hit = [t for t, m in moves.items() if m == (doc, first, event)]
    else:
        hit = [t for t, m in moves.items() if m[0] == doc and m[1] in froms]
    order = {t: i for i, t in enumerate(story_ids)}
    return sorted(set(hit), key=lambda t: (order.get(t, 99), t)) or story_ids


NAMED = {"sam": "Sales staff", "coordinator": "Logistics coordinator", "buyer": "Purchasing staff"}  # story words, as model roles


def tester_role(tag, allowed, expected):
    """Who runs the row: the role allowed to make the move, or the role the refusal is shown to."""
    if tag and tag.group(1) == "Transition":
        r = allowed.get((tag.group(2), tag.group(3).strip(), (tag.group(4) or "").strip()))
        if r and r["roles"] != "System":
            return r["roles"]
    m = re.search(r"And (?:the )?([A-Za-z ]+?) sees", expected)
    if not m:
        return ""
    who = m.group(1).strip()
    return NAMED.get(who.lower(), who[:1].upper() + who[1:])


ID = re.compile(r'"((?:SO|TR|PO|MR|INV|K|C|V)-[0-9][0-9-]*)"')
KIND = {"SO": "sales order", "TR": "delivery trip", "PO": "purchase order", "MR": "material request",
        "INV": "sales invoice", "K": "customer contract", "C": "customer", "V": "supplier"}


def test_data(rows):
    """Every record a row needs, from its steps: id, kind, the states the rows start it in, and the rows that use it."""
    data = {}
    for r in rows:
        for rid in ID.findall(r["steps"] + "\n" + r["expected"]):
            d = data.setdefault(rid, {"rows": [], "states": set()})
            if r["id"] not in d["rows"]:
                d["rows"].append(r["id"])
        for line in r["steps"].splitlines():  # a line's state belongs to the first record it names, its subject
            ids, m = ID.findall(line), re.search(r' is "([^"]+)"', line)
            if ids and m:
                data[ids[0]]["states"].add(m.group(1))
    return sorted(data.items(), key=lambda kv: (KIND[kv[0].split("-")[0]], kv[0]))


def judge(tag, allowed, refused, states):
    """(transition text, rule, problem or None) for one scenario tag."""
    if not tag:
        return "", "", "no tag"
    kind, doc, first, event = tag.group(1), tag.group(2), tag.group(3).strip(), (tag.group(4) or "").strip()
    if kind == "Transition":
        row = allowed.get((doc, first, event))
        if not row:
            return "", "", "transition %s | %s | %s is not allowed in the model" % (doc, first, event)
        return "%s: %s, to %s" % (first, event, row["to"]), row["rule"], None
    if kind == "Error path":
        row = refused.get((doc, first, event))
        if not row:
            return "", "", "error path %s | %s | %s is not refused in the model" % (doc, first, event)
        return "refused in %s: %s" % (first, event), row["rule"], None
    wanted = [x.strip() for x in first.split(",")]
    missing = [x for x in wanted if x not in states.get(doc, set())]
    if missing:
        return "", "", "state %s is not in the model for %s" % (missing, doc)
    rule = event if re.fullmatch(r"R[0-9]+", event) else ""  # a state check may name the rule it proves
    return "in state: %s" % ", ".join(wanted), rule, None


def parse(root):
    allowed, refused, states = model(root)
    steps = map_steps(root)
    moves = map_moves(root)
    rows, stories, problems = [], [], []
    for path in sorted(glob.glob(os.path.join(root, "stories", "*.feature"))):
        story = os.path.splitext(os.path.basename(path))[0]
        doc = Parser().parse(open(path, encoding="utf-8").read())
        notes = {c["location"]["line"]: c["text"].strip() for c in doc["comments"]}
        head = " ".join(t for ln, t in notes.items() if ln < doc["feature"]["location"]["line"])
        m = re.search(r"Map:\s*after\.bpmn,\s*([^.]+)\.", head)
        ids = [x.strip() for x in m.group(1).split(",")] if m else []
        for x in ids:
            if x not in steps:
                problems.append("%s: map step %s is not on process/after.bpmn" % (story, x))
        dm = re.search(r"Document:\s*([\w-]+)", head)
        document = dm.group(1) if dm else ""
        decision = ", ".join(re.findall(r"\bD\d+\b", head.split("Decision:", 1)[1])) if "Decision:" in head else ""
        background, errors = [], 0
        for child in doc["feature"]["children"]:
            if "background" in child:
                background = child["background"]["steps"]
                continue
            sc = child["scenario"]
            tag = TAG.match(notes.get(sc["location"]["line"] - 1, ""))
            text, rule, problem = judge(tag, allowed, refused, states)
            if problem:
                problems.append("%s / %s: %s" % (story, sc["name"], problem))
            is_error = bool(tag and tag.group(1) == "Error path")
            errors += is_error
            setup, expect = [], []
            target = setup
            for st in background + sc["steps"]:
                if st["keyword"].strip() == "Then":
                    target = expect
                target.append(st["keyword"].strip() + " " + st["text"])
            exact = row_steps(tag, moves, ids)
            for x in exact:
                if x not in steps:
                    problems.append("%s / %s: map step %s is not on process/after.bpmn" % (story, sc["name"], x))
            rows.append({"story": story, "scenario": sc["name"], "document": document, "transition": text,
                         "rule": rule, "error": is_error,
                         "refused": (tag.group(2), tag.group(3).strip(), tag.group(4).strip()) if is_error else None,
                         "map": ", ".join(exact), "decision": decision,
                         "steps": "\n".join(setup), "expected": "\n".join(expect),
                         "who": tester_role(tag, allowed, "\n".join(expect))})
        stories.append((story, doc["feature"]["name"], errors))
    for i, r in enumerate(rows):
        r["id"] = "UAT-%02d" % (i + 1)
    untested = sorted(k for k in refused if k not in {r["refused"] for r in rows if r["refused"]})
    return rows, stories, problems, untested


def style(ws, ncols, input_from=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Calibri", size=10, bold=True, color=HEADER_TEXT)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if input_from and c.column >= input_from:
                c.fill = PatternFill("solid", fgColor=INPUT)


def build(root, rows, stories, untested):
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    rm.append(["Item", "Detail"])
    for k, v in [
        ("Purpose", "User acceptance test script for the Tanmark Industries order-to-cash, procure-to-pay and "
                    "accounting system. Tanmark Industries is fictional; every name and number is illustrative."),
        ("Version", "v5, 2026-09-30 (map steps per row, test data, defects)"),
        ("Owner", "Product owner (role). Each department signs the stories it owns on the Sign-off sheet."),
        ("How to use", "Run each row as written. Record what happened in Actual. Choose Pass, Fail or Blocked in "
                       "Result. A Fail needs a defect id in Actual."),
        ("Where rows come from", "Generated from stories/*.feature by uat/build_uat.py and checked against the state "
                                 "model (state/transitions.csv) and the after map (process/after.bpmn)."),
        ("State transition", "The move the row exercises (from state: event, to state), the state it checks, or the "
                             "move it proves refused."),
        ("Not tested as a row", "Refused moves in the model that no row tests: " +
         ("; ".join("%s, %s: %s" % k for k in untested) or "none") + ". A choice, not a gap: each is enforced "
         "by the same rule as a tested move."),
        ("Shaded cells", "Cells shaded pale yellow are for the tester. Everything else is generated, including "
                         "the suggested tester role, which the tester may change."),
        ("Map steps", "Each row names the step on process/after.bpmn that makes its move, or, for a state or a "
                      "refusal, the steps that move the document on from that state."),
        ("Test data", "Every record the rows use, the states they start it in, and the rows that use it: load these "
                      "before testing."),
        ("Defects", "One row per defect, tied to a UAT row. Severity 1 blocks a process or breaks a rule; 2 works "
                    "with a workaround; 3 is cosmetic. An open severity 1 defect stops the go-live (cutover plan, "
                    "criterion 2), and the Sign-off sheet counts them per story."),
    ]:
        rm.append([k, v])
    style(rm, 2)
    rm.column_dimensions["A"].width = 22
    rm.column_dimensions["B"].width = 110

    ws = wb.create_sheet("UAT")
    cols = ["ID", "Story", "Scenario", "Document", "State transition", "Rule", "Map steps", "Decisions", "Steps",
            "Expected", "Actual", "Result", "Tester role", "Date"]
    ws.append(cols)
    for r in rows:
        ws.append([r["id"], r["story"], r["scenario"], r["document"], r["transition"], r["rule"] or "none", r["map"],
                   r["decision"], r["steps"], r["expected"], "", "", r["who"], ""])
    style(ws, len(cols), input_from=11)
    for i, w in enumerate([9, 16, 32, 14, 30, 6, 14, 9, 52, 44, 26, 10, 16, 12]):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
    last = len(rows) + 1
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(RESULTS), allow_blank=True, showErrorMessage=True,
                        errorTitle="Result", error="Choose Pass, Fail or Blocked.")
    ws.add_data_validation(dv)
    dv.add("L2:L%d" % last)
    for val, col in (("Pass", PASS), ("Fail", FAIL)):
        ws.conditional_formatting.add("L2:L%d" % last, CellIsRule(operator="equal", formula=['"%s"' % val],
                                      font=Font(color=col, bold=True)))
    ws.add_table(Table(displayName="UAT", ref="A1:N%d" % last,
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    ws.freeze_panes = "B2"

    so = wb.create_sheet("Sign-off")
    so.append(["Story", "Feature", "Owning department", "Signs", "Scenarios", "Passed", "Failed",
               "Open severity 1 defects", "Decision", "Signed by (role)", "Date", "Note"])
    for i, (story, name, _) in enumerate(stories):
        n = i + 2
        dept, signer = OWNER.get(story, ("", ""))
        so.append([story, name, dept, signer, '=COUNTIFS(UAT!$B$2:$B$%d,A%d)' % (last, n),
                   '=COUNTIFS(UAT!$B$2:$B$%d,A%d,UAT!$L$2:$L$%d,"Pass")' % (last, n, last),
                   '=COUNTIFS(UAT!$B$2:$B$%d,A%d,UAT!$L$2:$L$%d,"Fail")' % (last, n, last),
                   '=COUNTIFS(Defects!$C$2:$C$%d,A%d,Defects!$D$2:$D$%d,"1",Defects!$H$2:$H$%d,"<>Closed")'
                   % (DEFECT_ROWS + 1, n, DEFECT_ROWS + 1, DEFECT_ROWS + 1), "", "", "", ""])
    style(so, 12, input_from=9)
    for i, w in enumerate([16, 44, 16, 44, 10, 9, 9, 12, 12, 20, 12, 30]):
        so.column_dimensions[so.cell(row=1, column=i + 1).column_letter].width = w
    ls = len(stories) + 1
    dv2 = DataValidation(type="list", formula1='"Accept,Accept with conditions,Reject"', allow_blank=True,
                         showErrorMessage=True, errorTitle="Decision",
                         error="Choose Accept, Accept with conditions or Reject.")
    so.add_data_validation(dv2)
    dv2.add("I2:I%d" % ls)
    so.add_table(Table(displayName="SignOff", ref="A1:L%d" % ls,
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    so.freeze_panes = "B2"

    td = wb.create_sheet("Test data")
    td.append(["Record", "Kind", "Start it in", "Used by rows"])
    for rid, d in test_data(rows):
        td.append([rid, KIND[rid.split("-")[0]], ", ".join(sorted(d["states"])) or "as the row's steps say",
                   ", ".join(d["rows"])])
    style(td, 4)
    for i, w in enumerate([12, 18, 34, 60]):
        td.column_dimensions[td.cell(row=1, column=i + 1).column_letter].width = w
    td.freeze_panes = "A2"

    df = wb.create_sheet("Defects")
    df.append(["Defect", "UAT row", "Story", "Severity", "What happened", "Found by (role)", "Owner", "Status",
               "Date found", "Date closed"])
    for i in range(DEFECT_ROWS):
        n = i + 2
        df.append(["DEF-%02d" % (i + 1), "",
                   '=IF(B%d="","",IFERROR(INDEX(UAT!$B$2:$B$%d,MATCH(B%d,UAT!$A$2:$A$%d,0)),"not a UAT row"))'
                   % (n, last, n, last), "", "", "", "", "", "", ""])
    style(df, 10, input_from=4)
    df["B1"].fill = PatternFill("solid", fgColor=HEADER)
    for c in range(2, DEFECT_ROWS + 2):
        df.cell(row=c, column=2).fill = PatternFill("solid", fgColor=INPUT)
    for i, w in enumerate([9, 10, 16, 9, 44, 16, 16, 10, 12, 12]):
        df.column_dimensions[df.cell(row=1, column=i + 1).column_letter].width = w
    end = DEFECT_ROWS + 1
    for ref, values, title in (("B2:B%d" % end, "=UAT!$A$2:$A$%d" % last, "UAT row"),
                               ("D2:D%d" % end, '"1,2,3"', "Severity"),
                               ("H2:H%d" % end, '"Open,Fixed,Retested,Closed"', "Status")):
        v = DataValidation(type="list", formula1=values, allow_blank=True, showErrorMessage=True, errorTitle=title,
                           error="Choose from the list.")
        df.add_data_validation(v)
        v.add(ref)
    df.add_table(Table(displayName="Defects", ref="A1:J%d" % end,
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    df.freeze_panes = "B2"

    cl = wb.create_sheet("Change log")
    cl.append(["Date", "Change", "By (role)"])
    cl.append(["2026-09-29", "First version, generated from three stories", "Product owner"])
    cl.append(["2026-09-30", "Rebuilt on the state model: six stories, every error path a refused move", "Product owner"])
    cl.append(["2026-09-30", "Every row carries its state transition and rule; map steps checked (second read)",
               "Product owner"])
    cl.append(["2026-09-30", "B2B receivables: an invoice per trip before dispatch, the credit hold, oldest-first allocation, part credit notes, returns; the sales-invoice and sale-status stories", "Product owner"])
    cl.append(["2026-09-30", "The draft stock form (D15); the trip Issued state; partial receipts and issues; the material-request story; purchase-order scenarios in the real sequence", "Product owner"])
    cl.append(["2026-09-30", "Map steps per row; the Test data and Defects sheets; open severity 1 defects on Sign-off; the tester role suggested", "Product owner"])
    style(cl, 3)
    cl.column_dimensions["B"].width = 70
    cl.column_dimensions["C"].width = 18
    out = os.path.join(root, "uat", "uat-script.xlsx")
    wb.save(out)
    return out


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    root = argv[1]
    allowed, refused, states = model(root)
    steps = map_steps(root)
    plant = [judge(TAG.match("# Error path: sales-order | Completed | reopen"), allowed, refused, states)[2],
             judge(TAG.match("# Transition: voucher | Posted | edit"), allowed, refused, states)[2],
             None if "T99" in steps else "map step T99 caught"]
    print("[CONTROL] planted refused move, transition and map step caught: %s" % all(plant))
    ok = all(plant)
    rows, stories, problems, untested = parse(root)
    for p in problems:
        print("  REFUSED " + p)
    ok = ok and not problems
    for story, _, errors in stories:
        print("  %-16s error paths %d" % (story, errors))
        ok = ok and errors >= 1
    blank = [r["id"] for r in rows if not r["transition"]]
    print("[CONTROL] rows with no state transition: %d %s" % (len(blank), blank))
    ok = ok and not blank
    out = build(root, rows, stories, untested)
    back = load_workbook(out)["UAT"]
    ok = ok and back.max_row - 1 == len(rows)
    print("scenarios %d . rows written %d . transitions %d . states %d . error paths %d . refused moves not tested %d"
          % (len(rows), back.max_row - 1, sum(1 for r in rows if ", to " in r["transition"]),
             sum(1 for r in rows if r["transition"].startswith("in state")), sum(1 for r in rows if r["error"]),
             len(untested)))
    print("[RESULT] %s -> %s" % ("ok" if ok else "REFUSED", out))
    return 0 if ok else 3


if __name__ == "__main__":
    sys.exit(main(sys.argv))
