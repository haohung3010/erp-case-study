# -*- coding: utf-8 -*-
"""build_deck.py -- the ten-minute interview deck, drawn on the DesignOS slide grammar from the report's numbers.

PURPOSE
  The deck a BA or implementation panel sees. It follows the catalogue's interview-case pack (ten slides, in the
  pack's order) on the portfolio theme, light scheme, and it reads every number from docs/numbers.json, so it cannot
  disagree with the report or the README.

CONFIG
  argv[1]  the repository root. Needs python-pptx. Grammar and theme values are copied from the portfolio harness's
  projection of the catalogue (catalogue 0.10.1): canvas 1920 x 1080, 12 columns of 122 px, gutters 24, margins 96,
  title band 80-260, content 292-968, footer 1000-1048; Arial; type px caption 24, body 30, card title 36,
  sub heading 44, title 56, display 88. One px is 0.5 pt on a 13.333 in slide.

WORKFLOW
  1 the ten slides are declared below as data: template id, slots, speaker notes
  2 fill {{key}} from numbers.json; check every slot against its template's word limit; refuse an unfilled key, an em
    dash, or a slot over its limit
  3 draw each slide: title band, content by template, footer with the source line and the page number
  4 write deck/interview-deck.pptx; read it back and count slides and notes
  5 [CONTROL] the limit check is shown able to refuse a planted over-limit title before the real deck is checked
  rc 0 ok . 2 usage . 3 refused

Q&A
  Q: why draw rather than fill a template deck?   A: the catalogue ships rules, not a deck, and the pack's exported
                                                     deck is not in the harness; drawing from the grammar keeps every
                                                     position a grammar value.
  Q: where is the dark scheme?   A: not built. The light scheme is the report's, and a panel room usually projects it.

PITFALLS
  python-pptx does not shrink text to fit. The word limits are what keep text inside its region, so they are checked,
  and the rendered PDF is looked at by eye.

NOTES
  Tanmark Industries is fictional; every number is illustrative.
"""
import io
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

PX = 6350  # EMU per px: 13.333 in / 1920 px
COL, GUT, MARGIN = 122, 24, 96
C = {"ground": "FBFCFC", "surface": "F1F5F6", "ink": "16232B", "muted": "5B6B73", "line": "D5DDE0",
     "accent": "1D5C6B", "on_accent": "FFFFFF"}
FAMILY = [("1D5C6B", "93B1B8"), ("906915", "D3AF5F"), ("5C6E99", "9EACCB"), ("8F4E6B", "C4A3B2"), ("527A40", "97B48A")]
SIZE = {"caption": 24, "body": 30, "card_title": 36, "sub_heading": 44, "title": 56, "display": 88}
SOURCE = "Tanmark Industries is fictional. Numbers are illustrative, from the synthetic dataset in data/."
LIMITS = {  # words per slot, from the templates
    "T01": {"title": 8, "subtitle": 15}, "T03": {"title": 8}, "T04": {"statement": 15},
    "T05": {"title": 12, "lead": 60}, "T06": {"item": 12}, "T08": {"body": 50, "caption": 15},
    "T17": {"cell": 15}, "T18": {"step": 15}, "T11": {"heading": 5, "body": 40},
    "T24": {"step": 4, "branch": 2, "lane": 3},
}
T24_PACK = {"lanes": (2, 6), "steps": (3, 6)}  # interview-case pack settings for Before and After
T24_CONFLICTS = {(5, 7), (5, 8), (6, 7), (6, 8)}


def check_t24(i, s):
    problems = []
    lanes, nodes, edges = s["lanes"], s["nodes"], s["edges"]
    steps = max(n[2] for n in nodes)
    lo, hi = T24_PACK["lanes"]
    if not lo <= len(lanes) <= hi:
        problems.append("slide %d: %d lanes outside the pack's %d-%d" % (i, len(lanes), lo, hi))
    lo, hi = T24_PACK["steps"]
    if not lo <= steps <= hi:
        problems.append("slide %d: %d steps outside the pack's %d-%d" % (i, steps, lo, hi))
    if (len(lanes), steps) in T24_CONFLICTS:
        problems.append("slide %d: %d lanes with %d steps is a T24 conflict" % (i, len(lanes), steps))
    for ln in lanes:
        if words(ln) > LIMITS["T24"]["lane"]:
            problems.append("slide %d lane label over 3 words: %s" % (i, ln))
    seen = set()
    for nid, lane, col, label in nodes:
        if words(label) > LIMITS["T24"]["step"]:
            problems.append("slide %d step label over 4 words: %s" % (i, label))
        if not (0 <= lane < len(lanes) and 1 <= col <= steps) or (lane, col) in seen:
            problems.append("slide %d step %s is off the grid or shares a cell" % (i, nid))
        seen.add((lane, col))
    ids = {n[0] for n in nodes}
    forks = [a for a in ids if sum(1 for e in edges if e[0] == a) >= 2]
    if len(forks) > 2:
        problems.append("slide %d: %d forks, the limit is 2" % (i, len(forks)))
    for a, b, label in edges:
        if a not in ids or b not in ids:
            problems.append("slide %d: connector %s-%s names a missing step" % (i, a, b))
        if words(label) > LIMITS["T24"]["branch"]:
            problems.append("slide %d connector label over 2 words: %s" % (i, label))
    return problems
TITLE_WORDS = 14  # title band: two lines at 56 px

SLIDES = [
    {"t": "T01", "title": "From paper to one system",
     "subtitle": "How a tank, drum and container maker moved orders, purchasing and accounting onto one system",
     "eyebrow": "Product owner case study . Tanmark Industries, a fictional company",
     "notes": "This is a case study of a fictional manufacturer, built on real design patterns from order-processing and "
              "accounting systems, the kind of change I led. The company and every number are synthetic. One outcome is "
              "real, and I will say which."},
    {"t": "T04", "title": "The problem",
     "statement": "Nobody could see where an order was stuck.",
     "notes": "Orders on paper forms, stock checked by walking the shelves, purchases by email, invoices keyed by hand. "
              "Orders were cancelled after materials had gone to the workshop, and vouchers were edited and deleted after "
              "posting."},
    {"t": "T24", "title": "Before: paper, email and a hand-checked ledger",
     "lanes": ["Sales", "Logistics", "Purchasing", "Warehouse", "Production", "Accounting"],
     "nodes": [("a", 0, 1, "Write paper order"), ("b", 1, 2, "Walk the shelves"), ("c", 2, 3, "Email the supplier"),
               ("d", 4, 4, "Build from job card"), ("e", 3, 5, "Ship with paper note"),
               ("f", 5, 6, "Match invoices by hand")],
     "edges": [("a", "b", "form carried"), ("b", "c", "short"), ("b", "e", "in stock"), ("c", "d", ""),
               ("d", "e", ""), ("e", "f", "paper copies")],
     "notes": "Every hand-off is paper or email, so each one is a place an order waits unseen. About "
              "{{orders_per_month}} orders a month went this way. The full BPMN map is in the report."},
    {"t": "T18", "title": "What was done: states first, then everything that depends on them",
     "steps": [("Discovery", "When may each document still change, and who may change it?"),
               ("A state model", "Seven documents: states, allowed moves, refused moves"),
               ("Owners decide", "Approval, posting, credit and closure policy stay with the business"),
               ("Acceptance first", "Every story error path is a move the model refuses"),
               ("A pilot", "{{pilot_size}} people for {{pilot_weeks}} weeks, then a month in parallel (D2, D12)")],
     "notes": "The state model came before the screens. It has 44 allowed moves and 19 refused ones. The rejected big-bang "
              "option has a public reference case in Canada's Phoenix pay system: the Auditor General found its pilot "
              "had been cut and warnings did not travel up."},
    {"t": "T24", "title": "After, simplified: an order moves forward through named states",
     "lanes": ["Sales", "Logistics", "Production", "Accounting"],
     "nodes": [("a", 0, 1, "Confirm order; credit check"), ("c", 1, 2, "Check finished stock"),
               ("d", 2, 3, "Build with issued materials"), ("f", 3, 4, "Invoice the trip"),
               ("e", 1, 5, "Dispatch and deliver"), ("g", 3, 6, "Collect on terms")],
     "edges": [("a", "c", "locks sales"), ("c", "d", "to build"), ("c", "f", ""), ("d", "f", ""),
               ("f", "e", "invoice raised"), ("e", "g", "delivered")],
     "notes": "Sales confirms an order only if the customer is within its credit; handing it to logistics locks it for "
              "sales. Finished goods already in stock go straight to a trip. Each trip is invoiced before it leaves, and the customer pays on its contract's 30-day terms after "
              "delivery. Cancelling, purchasing with approval kept on paper, and the voucher that is posted and locked "
              "are in the full maps and the state model."},
    {"t": "T17", "title": "Every error path is a move the system refuses",
     "table": [["Story", "The refused move", "Signed off by"],
               ["Order lifecycle", "Sales edits an order that logistics has", "Sales"],
               ["Order closure", "Cancel an order already in production", "Sales, with logistics"],
               ["Delivery trip", "Dispatch a trip before its invoice is raised", "Logistics"],
               ["Voucher posting", "Edit a posted voucher, or delete any voucher", "Chief accountant"]],
     "summary": "47 scenarios in 8 stories. The build refuses a story whose error path is not in the state model.",
     "notes": "Sampled: one error path from four of the eight stories. The UAT script is generated from the stories, and the "
              "traceability matrix ties every test to a business objective."},
    {"t": "T11", "title": "Orders reached customers {{lead_change_pct}}% faster, and fewer were left unpaid",
     "columns": [("{{lead_after_days}} days, from {{lead_before_days}}",
                  "Median days from order to delivery, {{orders_after_n}} orders after the pilot against "
                  "{{orders_before_n}} before."),
                 ("{{overdue_after_pct}}% overdue, from {{overdue_before_pct}}%",
                  "Share of open receivables past due on 30-day terms, averaged by month."),
                 ("{{rework_after_pct}}% rework, from {{rework_before_pct}}%",
                  "Share of orders needing rework, averaged by month over the twelve months after pilot go-live."),
                 ("All {{pilot_size}} by week {{pilot_full_week}}",
                  "Pilot group members active in the system each week, over the {{pilot_weeks}}-week pilot.")],
     "notes": "These four figures are illustrative, from the synthetic dataset. The real outcome: the system this is "
              "modelled on is used by 50 to 100 people, and it replaced hand records, manual warehouse checks and manual "
              "order processing. No real timings exist, so I do not claim any."},
    {"t": "T06", "title": "What went wrong, and what changed because of it",
     "items": ["Collection stayed on paper in the pilot; overdue stayed at {{overdue_pilot_pct}}%",
               "Adoption stuck at {{pilot_plateau_active}} of {{pilot_size}} until the night shift was trained",
               "Three decided rules were missing from the first build; UAT caught them",
               "Rework peaked at {{rework_golive_pct}}% in the pilot's first month"],
     "notes": "Each problem became a rule in the rollout plan. The decided-versus-built register is why UAT caught the "
              "missing rules: every decided rule had its own test. Next time I would put collection in the pilot."},
    {"t": "T05", "title": "For your organisation: model the states before the screens",
     "lead": "Ask of every document when it may still change, and who may change it. Write the answers as states, "
             "moves and refusals, and test each refusal. Leave policy with its owners: approval, posting, credit and "
             "closure belong to the business. A decided rule is not a built rule until a test proves it.",
     "notes": "Tie this to their process: which of their documents can still be changed after it matters?"},
    {"t": "T03", "title": "The full case, in one repository",
     "cta": "State model, maps, stories, UAT, RACI and thirteen working papers",
     "notes": "Everything in the deck is in the repository, generated from one model and one dataset, so the report and "
              "the deck agree."},
]


def fill(obj, nums):
    if isinstance(obj, str):
        return re.sub(r"\{\{(\w+)\}\}", lambda m: nums.get(m.group(1), m.group(0)), obj)
    if isinstance(obj, (list, tuple)):
        return type(obj)(fill(x, nums) for x in obj)
    if isinstance(obj, dict):
        return {k: fill(v, nums) for k, v in obj.items()}
    return obj


def words(s):
    return len(s.split())


def check(slides):
    problems = []
    for i, s in enumerate(slides, 1):
        lim = LIMITS[s["t"]]
        if s["t"] == "T24":
            problems += check_t24(i, s)
        blob = json.dumps(s, ensure_ascii=False)
        if "{{" in blob:
            problems.append("slide %d: unfilled %s" % (i, re.findall(r"\{\{\w+\}\}", blob)))
        if "—" in blob:
            problems.append("slide %d: em dash" % i)
        if words(s["title"]) > lim.get("title", TITLE_WORDS):
            problems.append("slide %d title %d words > %d" % (i, words(s["title"]), lim.get("title", TITLE_WORDS)))
        for slot in ("subtitle", "statement", "lead", "body", "caption"):
            if slot in s and words(s[slot]) > lim[slot]:
                problems.append("slide %d %s %d words > %d" % (i, slot, words(s[slot]), lim[slot]))
        for it in s.get("items", []):
            if words(it) > lim["item"]:
                problems.append("slide %d item %d words > %d: %s" % (i, words(it), lim["item"], it[:40]))
        for h, b in s.get("steps", []):
            if words(h) + words(b) > lim["step"]:
                problems.append("slide %d step %d words > %d" % (i, words(h) + words(b), lim["step"]))
        for h, b in s.get("columns", []):
            if words(h) > lim["heading"] or words(b) > lim["body"]:
                problems.append("slide %d column over limit: %s" % (i, h))
        for row in s.get("table", []):
            for cell in row:
                if words(cell) > lim["cell"]:
                    problems.append("slide %d cell %d words > %d" % (i, words(cell), lim["cell"]))
    return problems


def col_x(a):
    return MARGIN + (a - 1) * (COL + GUT)


def span_w(a, b):
    return (b - a + 1) * COL + (b - a) * GUT


def box(slide, x, y, w, h, text, size="body", color="ink", bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        fill=None):
    shp = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX)) if not fill else \
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(12 * PX if fill else 0))
    lines = text if isinstance(text, list) else [text]
    for j, line in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = "Arial"
        r.font.size = Pt(SIZE[size] / 2)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(C.get(color, color))
    return shp


def rule(slide, x1, y1, x2, y2, color="line", width=2):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1 * PX), Emu(y1 * PX), Emu(x2 * PX), Emu(y2 * PX))
    ln.line.color.rgb = RGBColor.from_string(C.get(color, color))
    ln.line.width = Pt(width / 2)
    return ln


def arrow_line(slide, pts, color="ink"):
    """A polyline of straight connectors; the last segment carries an arrowhead."""
    from pptx.oxml.ns import qn
    for j in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[j], pts[j + 1]
        ln = rule(slide, x1, y1, x2, y2, color, 3)
        if j == len(pts) - 2:
            tail = ln.line._get_or_add_ln().makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
            ln.line._get_or_add_ln().append(tail)


def swimlane(sl, s):
    lanes, nodes, edges = s["lanes"], s["nodes"], s["edges"]
    steps = max(n[2] for n in nodes)
    top, height = 296, 668
    lane_h = height / len(lanes)
    x0, width = col_x(3), span_w(3, 12)
    step_w = width / steps
    bw, bh = step_w - 60, min(lane_h - 24, 100)  # up to three lines of caption text; the gap holds the connectors
    for k, name in enumerate(lanes):
        y = top + k * lane_h
        if k % 2 == 0:
            band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(MARGIN * PX), Emu(int(y * PX)), Emu(1728 * PX),
                                       Emu(int(lane_h * PX)))
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor.from_string(C["surface"])
            band.line.fill.background()
            band.shadow.inherit = False
        box(sl, MARGIN + 16, y, span_w(1, 2) - 16, lane_h, name, "caption", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rule(sl, MARGIN, top + height, MARGIN + 1728, top + height)
    place = {}
    for nid, lane, col, label in nodes:
        x = x0 + (col - 1) * step_w + (step_w - bw) / 2
        y = top + lane * lane_h + (lane_h - bh) / 2
        place[nid] = (x, y, bw, bh)
        shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(bw * PX)),
                                  Emu(int(bh * PX)))
        shp.adjustments[0] = 0.08
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        shp.line.color.rgb = RGBColor.from_string(C["accent"])
        shp.line.width = Pt(1.5)
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Emu(8 * PX))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = "Arial"
        r.font.size = Pt(SIZE["caption"] / 2)
        r.font.color.rgb = RGBColor.from_string(C["ink"])
    seen_into = {}  # connectors entering one step from other lanes each get their own vertical run
    for a, b, label in edges:
        ax, ay, aw, ah = place[a]
        bx, by, bw_, bh_ = place[b]
        if abs(ax - bx) < 1:  # same column: straight up or down
            sx = ax + aw / 2
            pts = [(sx, ay), (sx, by + bh_)] if by < ay else [(sx, ay + ah), (sx, by)]
            lx, ly = sx + 10, (pts[0][1] + pts[1][1]) / 2 - 16
        elif abs(ay - by) < 1:  # same lane
            pts = [(ax + aw, ay + ah / 2), (bx, by + bh_ / 2)]
            lx, ly = (ax + aw + bx) / 2 - 75, ay - 34  # above the boxes, centred on the gap
        else:  # along the source lane, down or up beside the target, then in
            k = seen_into.get(b, 0)
            seen_into[b] = k + 1
            xm = bx - 18 - 14 * k
            pts = [(ax + aw, ay + ah / 2), (xm, ay + ah / 2), (xm, by + bh_ / 2), (bx, by + bh_ / 2)]
            lx, ly = xm - 150, (ay + ah / 2 + by + bh_ / 2) / 2 - 16
        arrow_line(sl, pts, "muted")
        if label:
            same_lane = abs(ay - by) < 1 and abs(ax - bx) >= 1
            box(sl, lx, ly, 150, 34, label, "caption", "muted",
                align=PP_ALIGN.CENTER if same_lane else (PP_ALIGN.RIGHT if lx < pts[-1][0] - 100 else PP_ALIGN.LEFT))


def picture(slide, root, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(os.path.join(root, path)).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(os.path.join(root, path), Emu(int((x + (w - pw) / 2) * PX)),
                             Emu(int((y + (h - ph) / 2) * PX)), Emu(int(pw * PX)), Emu(int(ph * PX)))


def draw(prs, root, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor.from_string(C["ground"])
    t = s["t"]
    if t == "T01":
        box(sl, col_x(1), 300, span_w(1, 10), 40, s["eyebrow"], "caption", "muted")
        box(sl, col_x(1), 360, span_w(1, 10), 220, s["title"], "display", bold=True)
        box(sl, col_x(1), 600, span_w(1, 10), 140, s["subtitle"], "sub_heading", "accent")
        rule(sl, col_x(1), 780, col_x(1) + span_w(1, 3), 780, "accent", 8)
    elif t == "T03":
        box(sl, col_x(2), 380, span_w(2, 11), 120, s["title"], "display", bold=True)
        box(sl, col_x(2), 560, span_w(2, 11), 90, s["cta"], "sub_heading", "on_accent", fill=C["accent"],
            anchor=MSO_ANCHOR.MIDDLE)
    else:
        box(sl, col_x(1), 110, span_w(1, 12), 150, s["title"], "title", bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if t == "T04":
        box(sl, col_x(1), 292, span_w(1, 10), 676, s["statement"], "display", "accent", bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    elif t == "T08":
        box(sl, col_x(1), 292, span_w(1, 6) - 24, 676, s["body"], "body")
        picture(sl, root, s["media"], col_x(7), 292, span_w(7, 12), 600)
        box(sl, col_x(7), 912, span_w(7, 12), 48, s["caption"], "caption", "muted")
    elif t == "T24":
        swimlane(sl, s)
    elif t == "T18":
        k = len(s["steps"])
        w = (1728 - (k - 1) * 48) / k
        for j, (h, b) in enumerate(s["steps"]):
            x = MARGIN + j * (w + 48)
            mark, soft = FAMILY[j % 5]
            box(sl, x, 340, 72, 72, str(j + 1), "card_title", "on_accent", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                fill=mark)
            if j < k - 1:
                rule(sl, x + 88, 376, x + w + 32, 376, "muted", 3).line.end_arrowhead = None
            box(sl, x, 440, w, 100, h, "card_title", bold=True)
            box(sl, x, 540, w, 380, b, "body")
    elif t == "T17":
        rows = s["table"]
        widths = [440, 900, 388]
        y = 320
        for ri, row in enumerate(rows):
            x = MARGIN
            hgt = 72 if ri == 0 else 120
            for ci, cell in enumerate(row):
                box(sl, x, y, widths[ci] - 16, hgt, cell, "body", bold=(ri == 0), anchor=MSO_ANCHOR.MIDDLE,
                    fill=C["surface"] if ri == 0 else None)
                x += widths[ci]
            y += hgt
            rule(sl, MARGIN, y, MARGIN + 1728, y)
        box(sl, MARGIN, y + 32, 1728, 80, s["summary"], "caption", "muted")
    elif t == "T11":
        for j, (h, b) in enumerate(s["columns"]):
            a = 1 + 4 * j
            mark, soft = FAMILY[j]
            box(sl, col_x(a), 330, span_w(a, a + 3), 150, h, "sub_heading", mark, True, anchor=MSO_ANCHOR.BOTTOM)
            rule(sl, col_x(a), 500, col_x(a) + span_w(a, a + 3), 500, mark, 6)
            box(sl, col_x(a), 530, span_w(a, a + 3), 400, b, "body")
    elif t == "T06":
        y = 330
        for it in s["items"]:
            rule(sl, col_x(1), y + 22, col_x(1) + 24, y + 22, "accent", 12)
            box(sl, col_x(1) + 48, y, span_w(1, 8) + 400, 120, it, "body")
            y += 150
    elif t == "T05":
        box(sl, col_x(1), 330, span_w(1, 8), 600, s["lead"], "sub_heading")
    if t not in ("T01", "T03"):
        box(sl, MARGIN, 1000, 1400, 48, SOURCE, "caption", "muted", anchor=MSO_ANCHOR.MIDDLE)
        box(sl, 1624, 1000, 200, 48, "%d / %d" % (n, total), "caption", "muted", align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE)
    sl.notes_slide.notes_text_frame.text = s["notes"]


NUM = re.compile(r"(?<![\w.])\d+(?:[.,]\d+)?%?")


def parity(root, prs):
    """Every number on a slide or in its notes must appear in the report's text. None when the report is absent."""
    pdf = os.path.join(root, "docs", "case-study-report.pdf")
    if not os.path.exists(pdf):
        return None
    from pypdf import PdfReader
    rep = set(NUM.findall(" ".join(p.extract_text() or "" for p in PdfReader(pdf).pages)))
    ignore = {str(k) for k in range(1, 11)}  # page and step numbers
    assert set(NUM.findall("a pilot of 37 people")) - rep - ignore == {"37"}, "[CONTROL] parity check cannot fire"
    missing = {}
    for i, s in enumerate(prs.slides, 1):
        text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        text += " " + s.notes_slide.notes_text_frame.text
        for n in set(NUM.findall(text)) - rep - ignore:
            missing.setdefault(n, []).append(i)
    return missing


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    root = argv[1]
    nums = json.load(io.open(os.path.join(root, "docs", "numbers.json"), encoding="utf-8"))
    planted = check([{"t": "T01", "title": "one two three four five six seven eight nine", "subtitle": "x",
                      "notes": ""}])
    ok = len(planted) == 1
    print("[CONTROL] planted nine-word cover title -> %s" % planted)
    slides = fill(SLIDES, nums)
    problems = check(slides)
    for p in problems:
        print("  REFUSED " + p)
    if problems or not ok:
        print("[RESULT] REFUSED")
        return 3
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    for i, s in enumerate(slides, 1):
        draw(prs, root, s, i, len(slides))
    out = os.path.join(root, "deck", "interview-deck.pptx")
    prs.save(out)
    back = Presentation(out)
    notes = sum(1 for s in back.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    print("slides %d . with speaker notes %d . templates %s" % (len(back.slides), notes, [s["t"] for s in slides]))
    ok = len(back.slides) == 10 and notes == 10
    # The report and the deck make the same claims with the same numbers.
    missing = parity(root, back)
    if missing is None:
        print("[BOUND] docs/case-study-report.pdf not built; the deck-to-report number check did not run")
    else:
        print("deck numbers not in the report: %s" % (missing or "none"))
        ok = ok and not missing
    print("[RESULT] %s -> %s" % ("ok" if ok else "REFUSED", out))
    return 0 if ok else 3


if __name__ == "__main__":
    sys.exit(main(sys.argv))
